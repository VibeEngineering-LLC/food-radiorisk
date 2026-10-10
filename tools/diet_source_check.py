"""Проверка числовых значений из diet.json по исходным файлам (Excel и Markdown)."""
import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

import xlrd
import openpyxl


def load_json(path: Path) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def col_to_index(col_str: str) -> int:
    index = 0
    for char in col_str:
        index = index * 26 + (ord(char.upper()) - ord('A') + 1)
    return index - 1


def read_cell(path: Path, sheet_name: str, cell_addr: str):
    if path.suffix.lower() == '.xls':
        if not hasattr(read_cell, 'xls_cache'):
            read_cell.xls_cache = {}
        if str(path) not in read_cell.xls_cache:
            read_cell.xls_cache[str(path)] = xlrd.open_workbook(str(path))
        wb = read_cell.xls_cache[str(path)]
        try:
            sheet = wb.sheet_by_name(sheet_name)
        except xlrd.biffh.XLRDError:
            return None
        match = re.match(r'^([A-Za-z]+)(\d+)$', cell_addr)
        if not match:
            return None
        col_str, row_str = match.groups()
        col_idx = col_to_index(col_str)
        row_idx = int(row_str) - 1
        return sheet.cell_value(row_idx, col_idx)
    elif path.suffix.lower() == '.xlsx':
        if not hasattr(read_cell, 'xlsx_cache'):
            read_cell.xlsx_cache = {}
        if str(path) not in read_cell.xlsx_cache:
            read_cell.xlsx_cache[str(path)] = openpyxl.load_workbook(str(path), data_only=True)
        wb = read_cell.xlsx_cache[str(path)]
        if sheet_name not in wb.sheetnames:
            return None
        ws = wb[sheet_name]
        return ws[cell_addr].value
    return None


def numbers_in(line: str):
    line = re.sub(r'(?<=\d),(?=\d)', '.', line)
    matches = re.findall(r"\d+(?:\.\d+)?", line)
    return [float(m) for m in matches]


def check_record(record: dict, sources: dict, problems: list, checked_count: list):
    rid = record.get('id', 'unknown')
    source_key = record.get('source')
    if not source_key or source_key not in sources:
        problems.append(f"{rid}: source '{source_key}' not found in sources")
        return

    src_info = sources[source_key]
    local_path_str = src_info.get('local')
    local_path = Path(local_path_str) if local_path_str else None

    wanted_value = record.get('src_value', record.get('value'))
    if wanted_value is None:
        problems.append(f"{rid}: no value or src_value")
        return

    cell_ref = record.get('cell')
    loc = record.get('loc', '')
    file_path = record.get('file')

    has_cell_check = bool(cell_ref)
    has_md_check = bool(re.search(r'([A-Za-z0-9_.\-]+\.(?:md|html)):(\d+)', loc))

    if not has_cell_check and not has_md_check:
        return

    passed = True
    checked_this_record = False

    if has_cell_check:
        parts = cell_ref.rsplit('!', 1)
        if len(parts) != 2:
            problems.append(f"{rid}: invalid cell reference '{cell_ref}'")
            passed = False
        else:
            sheet_name, cell_addr = parts
            val = read_cell(local_path, sheet_name, cell_addr) if local_path and local_path.exists() else None
            if not isinstance(val, (int, float)) or isinstance(val, bool):
                problems.append(f"{rid}: cell {cell_ref} is not a number (got {val}; file {local_path})")
                passed = False
            else:
                if abs(round(val, 3) - round(wanted_value, 3)) > 1e-9:
                    problems.append(f"{rid}: cell {cell_ref} value {val} != expected {wanted_value}")
                    passed = False
                else:
                    checked_this_record = True

    if has_md_check:
        match = re.search(r'([A-Za-z0-9_.\-]+\.(?:md|html)):(\d+)', loc)
        if match:
            md_file_name = match.group(1)
            line_num = int(match.group(2))
            
            if file_path:
                md_path = Path(file_path)
            else:
                md_path = (local_path.parent if local_path else Path('.')) / md_file_name

            if not md_path.exists():
                problems.append(f"{rid}: markdown file not found: {md_path}")
                passed = False
            else:
                try:
                    with open(md_path, 'r', encoding='utf-8') as f:
                        lines = f.read().split('\n')  # не splitlines(): форм-фид и \x1c из PDF сдвигают нумерацию строк относительно grep -n и редакторов
                    if line_num > len(lines) or line_num < 1:
                        problems.append(f"{rid}: line {line_num} out of range in {md_path}")
                        passed = False
                    else:
                        line_content = lines[line_num - 1]
                        nums = numbers_in(line_content)
                        found = False
                        for n in nums:
                            if abs(n - wanted_value) <= 1e-9:
                                found = True
                                break
                        if not found:
                            problems.append(f"{rid}: value {wanted_value} not found in line {line_num} of {md_path}")
                            passed = False
                        else:
                            checked_this_record = True
                except Exception as e:
                    problems.append(f"{rid}: error reading {md_path}: {e}")
                    passed = False

    if passed and checked_this_record:
        checked_count[0] += 1


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    
    parser = argparse.ArgumentParser(description="Verify diet.json values against sources")
    script_dir = Path(__file__).resolve().parent
    repo_root = script_dir.parent
    
    default_diet = repo_root / "public" / "data" / "diet.json"
    default_sources = repo_root / "public" / "data" / "sources.json"
    
    parser.add_argument('--diet', type=Path, default=default_diet)
    parser.add_argument('--sources', type=Path, default=default_sources)
    
    args = parser.parse_args()
    
    diet_path = args.diet
    sources_path = args.sources
    
    if not diet_path.exists():
        print(f"Error: diet file not found: {diet_path}")
        sys.exit(1)
    if not sources_path.exists():
        print(f"Error: sources file not found: {sources_path}")
        sys.exit(1)
        
    diet_data = load_json(diet_path)
    sources_data = load_json(sources_path)
    
    records = diet_data.get('records', [])
    sources = sources_data.get('sources', {})
    
    problems = []
    checked_count = [0]
    
    # Check 1: SHA256 for sources used by value records
    used_sources = set()
    for rec in records:
        if rec.get('kind') == 'value':
            src_key = rec.get('source')
            if src_key:
                used_sources.add(src_key)
                
    for src_key in used_sources:
        if src_key not in sources:
            continue
        src_info = sources[src_key]
        sha256_expected = src_info.get('sha256')
        local_path_str = src_info.get('local')
        
        if sha256_expected and local_path_str:
            local_path = Path(local_path_str)
            if not local_path.exists():
                problems.append(f"{src_key}: file not found for sha256 check: {local_path}")
            else:
                try:
                    sha256_hash = hashlib.sha256()
                    with open(local_path, "rb") as f:
                        for chunk in iter(lambda: f.read(8192), b""):
                            sha256_hash.update(chunk)
                    sha256_actual = sha256_hash.hexdigest()
                    if sha256_actual != sha256_expected:
                        problems.append(f"{src_key}: sha256 mismatch (expected {sha256_expected}, got {sha256_actual})")
                except Exception as e:
                    problems.append(f"{src_key}: error calculating sha256: {e}")

    # Check 2 & 3: Value records
    for rec in records:
        if rec.get('kind') != 'value':
            continue
        check_record(rec, sources, problems, checked_count)
        
    for p in problems:
        print(p)
        
    print(f"расхождений: {len(problems)}, сверено: {checked_count[0]}")
    
    return 0 if len(problems) == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
